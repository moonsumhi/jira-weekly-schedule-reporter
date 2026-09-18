"""Readable SR workbooks and a portable bundle of their original attachments."""
from __future__ import annotations

import io
import math
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from tempfile import SpooledTemporaryFile
from urllib.parse import unquote, urlsplit
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import html
from markdown_it import MarkdownIt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from PIL import Image, ImageOps, UnidentifiedImageError

from app.models.sr.service_request import REQUEST_TYPE_LABEL
from app.services.sr.export_labels import DETAIL_LABELS, DETAIL_OPTIONS, EDITOR_FIELDS, REQUEST_FIELDS
from app.utils.time import KST

UPLOAD_ROOT = Path('/app/uploads')
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
NAVY, BLUE, PALE, INK, MUTED = '243C5A', '365F91', 'EEF3F8', '243447', '64748B'
LINE = Side(style='thin', color='DCE4ED')
COLUMNS, COLUMN_WIDTH = 12, 9


def clean(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value if value is not None else ''))


def safe_filename(value, fallback='첨부파일'):
    value = unicodedata.normalize('NFC', clean(value))
    value = re.sub(r'[<>:"/\\|?*\r\n\t]', '_', value).strip(' .')
    if not value:
        return fallback
    stem, suffix = str(Path(value).stem), Path(value).suffix[:16]
    return stem[:120] + suffix


def date_text(value, *, time=False):
    if not value:
        return '—'
    try:
        dt = value if isinstance(value, datetime) else datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        if dt.tzinfo is None:
            # Request datetime-local inputs are stored without an offset in KST.
            dt = dt.replace(tzinfo=KST)
        return dt.astimezone(KST).strftime('%Y-%m-%d %H:%M' if time else '%Y-%m-%d')
    except (TypeError, ValueError):
        return clean(value)


def rich_root(value):
    rendered = MarkdownIt('commonmark', {'html': True, 'breaks': True}).enable('table').render(clean(value))
    return html.fragment_fromstring(rendered, create_parent='div')


def plain_text(value):
    root = rich_root(value)
    for node in root.xpath('.//script|.//style'):
        node.drop_tree()
    for node in root.iterdescendants():
        if node.tag == 'img':
            node.tail = (' [이미지] ' + (node.tail or ''))
        elif node.tag == 'br':
            node.tail = '\n' + (node.tail or '').lstrip('\n')
        elif node.tag == 'li':
            node.text = '• ' + (node.text or '')
        if node.tag in {'p', 'div', 'li', 'tr', 'h1', 'h2', 'h3', 'h4', 'blockquote', 'pre'}:
            node.tail = '\n' + (node.tail or '')
        elif node.tag in {'td', 'th'}:
            node.tail = '  |  ' + (node.tail or '')
    return re.sub(r'\n[ \t]*\n[ \t]*\n+', '\n\n', root.text_content()).strip()


def wrapped_lines(value, width=102):
    """Estimate Excel line wrapping including double-width Korean characters."""
    lines = []
    for raw in clean(value).split('\n'):
        line, used = '', 0
        for char in raw:
            size = 2 if unicodedata.east_asian_width(char) in ('W', 'F') else 1
            if used + size > width and line:
                lines.append(line)
                line, used = '', 0
            line += char
            used += size
        lines.append(line)
    return lines or ['—']


@dataclass
class Attachment:
    number: int
    name: str
    path: Path | None
    sources: list[str] = field(default_factory=list)
    image: bool = False
    warning: str = ''
    preview: bytes | None = None
    size: int = 0
    bundled: bool = False

    @property
    def archive_name(self):
        return f'첨부파일/{self.number:03d}_{safe_filename(self.name)}'


class Attachments:
    def __init__(self):
        self.items = []
        self.by_key = {}

    def add(self, source, attachment=None, url=None):
        attachment = attachment or {}
        raw_url = url or attachment.get('url', '')
        try:
            parsed = urlsplit(raw_url)
        except ValueError:
            parsed = urlsplit('')
        relative = unquote(parsed.path)
        match = re.fullmatch(r'/(?:api/)?uploads/(sr|pm)/([\w.\-]+)', relative) if not parsed.netloc and not parsed.scheme else None
        file_id = attachment.get('file_id', '')
        if match:
            relative = f'{match[1]}/{match[2]}'
        elif file_id and re.fullmatch(r'[\w.\-]+', file_id) and file_id not in ('.', '..'):
            relative = f'sr/{file_id}'
        else:
            relative = ''
        path = (UPLOAD_ROOT / relative).resolve() if relative else None
        if path and (not path.is_relative_to(UPLOAD_ROOT.resolve()) or not path.is_file()):
            path = None
        key = relative or raw_url or f'unknown-{len(self.items)}'
        if key in self.by_key:
            item = self.by_key[key]
            if source not in item.sources:
                item.sources.append(source)
            return item
        name = attachment.get('original_name') or Path(relative or parsed.path).name or '이미지'
        is_image = bool(url) or attachment.get('content_type', '').startswith('image/') or Path(name).suffix.lower() in ('.png', '.jpg', '.jpeg', '.gif', '.webp', '.bmp')
        item = Attachment(len(self.items) + 1, name, path, [source], is_image)
        if path is None:
            item.warning = '외부 이미지는 포함되지 않았습니다.' if parsed.netloc else '원본 파일을 찾을 수 없습니다.'
        self.by_key[key] = item
        self.items.append(item)
        return item

    def scan(self, value, source):
        if isinstance(value, dict):
            for key, text in value.items():
                self.scan(text, f'{source} · {DETAIL_LABELS.get(key, key)}')
        elif isinstance(value, list):
            for index, text in enumerate(value, 1):
                self.scan(text, f'{source} {index}')
        elif isinstance(value, str):
            for node in rich_root(value).xpath('.//img[@src]'):
                self.add(source, url=node.get('src'))

    def prepare(self, archive=None):
        for item in self.items:
            if item.path is None:
                continue
            try:
                item.size = item.path.stat().st_size
                # Original files are streamed into the temporary ZIP rather than read all at once.
                if archive is not None:
                    archive.write(item.path, item.archive_name)
                    item.bundled = True
            except OSError:
                item.warning = '원본 파일을 읽을 수 없습니다.'
                continue
            if not item.image:
                continue
            try:
                with Image.open(item.path) as source:
                    source.seek(0)
                    preview = ImageOps.exif_transpose(source)
                    preview.thumbnail((1600, 1600))
                    preview = preview.convert('RGBA')
                    buf = io.BytesIO()
                    preview.save(buf, format='PNG')
                    item.preview = buf.getvalue()
            except (OSError, ValueError, UnidentifiedImageError, Image.DecompressionBombError):
                item.warning = '이미지 미리보기를 만들 수 없습니다. 첨부한 원본 파일을 확인해 주세요.'


class ReportSheet:
    def __init__(self, wb, name, sr_no, title, subtitle):
        self.ws = wb.create_sheet(name)
        self.row = 1
        self.rendered_images = set()
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = BLUE
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4, header=.15, footer=.2)
        ws.oddFooter.center.text = '페이지 &P / &N'
        ws.oddFooter.right.text = clean(sr_no).replace('&', '&&')
        ws.print_options.horizontalCentered = True
        ws.freeze_panes = 'A5'
        ws.print_title_rows = '1:3'
        for col in range(1, COLUMNS + 1):
            ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTH
        self.band(f'{sr_no}  ·  {name}', bg=NAVY, fg='FFFFFF', size=12, height=30)
        self.band(title, size=19, height=max(42, len(wrapped_lines(title, 70)) * 25))
        self.band(subtitle, fg=MUTED, size=10, height=24)
        self.row += 1

    def cell(self, row, first, last, value, *, bg=None, fg=INK, size=11, bold=False):
        ws = self.ws
        if last > first:
            ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
        cell = ws.cell(row, first, clean(value))
        cell.data_type = 's'  # Uploaded text beginning with '=' must remain text, never a formula.
        cell.font = Font(name='맑은 고딕', size=size, color=fg, bold=bold)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        for col in range(first, last + 1):
            c = ws.cell(row, col)
            if bg:
                c.fill = PatternFill('solid', fgColor=bg)
            c.border = Border(bottom=LINE)
        return cell

    def band(self, value, *, bg=None, fg=INK, size=11, height=28):
        cell = self.cell(self.row, 1, COLUMNS, value, bg=bg, fg=fg, size=size, bold=True)
        self.ws.row_dimensions[self.row].height = max(height, len(wrapped_lines(value, int(102 * 11 / size))) * (size + 6) + 12)
        self.row += 1
        return cell

    def section(self, title):
        self.row += 1
        self.band(title, bg=PALE, fg=BLUE, height=30)

    def text(self, value):
        lines = wrapped_lines(value or '—')
        for start in range(0, len(lines), 12):
            part = lines[start:start + 12]
            self.cell(self.row, 1, COLUMNS, '\n'.join(part))
            self.ws.row_dimensions[self.row].height = max(30, len(part) * 17 + 14)
            self.row += 1

    def detail(self, label, value, attachments):
        if value is None or value == '':
            return
        self.section(label)
        self.rich_text(value, attachments, label)

    def picture(self, item):
        self.rendered_images.add(item.number)
        if item.preview is None:
            self.text(f'{item.name} · {item.warning or "이미지를 표시할 수 없습니다."}')
            return
        picture = ExcelImage(io.BytesIO(item.preview))
        ratio = min(1, 740 / picture.width, 600 / picture.height)
        picture.width, picture.height = picture.width * ratio, picture.height * ratio
        self.ws.add_image(picture, f'A{self.row}')
        count = math.ceil(picture.height / 20) + 1
        for row in range(self.row, self.row + count):
            self.ws.row_dimensions[row].height = 15
        self.row += count

    def rich_text(self, value, attachments, label):
        """Keep text and images in document order on the same worksheet."""
        root = rich_root(value)
        pending = []

        def flush():
            text = re.sub(r'\n[ \t]*\n[ \t]*\n+', '\n\n', ''.join(pending)).strip()
            if text:
                self.text(text)
            pending.clear()

        def walk(node):
            if node.tag in {'script', 'style'} or not isinstance(node.tag, str):
                return
            if node.tag == 'img':
                flush()
                self.picture(attachments.add(label, url=node.get('src', '')))
                return
            if node.tag == 'table':
                flush()
                rows = node.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
                cells = [row.xpath('./th|./td') for row in rows]
                if cells:
                    values = [[plain_text(html.tostring(cell, encoding='unicode')) for cell in row] for row in cells]
                    width = max(map(len, values))
                    has_header = any(cell.tag == 'th' for cell in cells[0])
                    headers = values[0] if has_header else [''] * width
                    body = values[1:] if has_header else values
                    for start in range(0, width, 6):
                        end = min(start + 6, width)
                        self.table([(headers + [''] * width)[i] for i in range(start, end)],
                                   [[(row + [''] * width)[i] for i in range(start, end)] for row in body])
                    for img in node.xpath('.//img[@src]'):
                        self.picture(attachments.add(label, url=img.get('src')))
                return
            if node.tag == 'br':
                pending.append('\n')
            elif node.tag == 'li':
                pending.append('• ')
            if node.text:
                pending.append(node.text)
            for child in node:
                walk(child)
                if child.tail:
                    pending.append(child.tail.lstrip('\n') if child.tag == 'br' else child.tail)
            if node.tag in {'p', 'div', 'li', 'h1', 'h2', 'h3', 'h4', 'blockquote', 'pre'}:
                pending.append('\n')

        walk(root)
        flush()

    def table(self, headers, rows, spans=None):
        spans = list(spans) if spans else [COLUMNS // len(headers)] * len(headers)
        spans[-1] += COLUMNS - sum(spans)
        first, columns = 1, []
        for span, header in zip(spans, headers):
            columns.append((first, first + span - 1, span * COLUMN_WIDTH - 2))
            self.cell(self.row, first, first + span - 1, header, bg=BLUE, fg='FFFFFF', bold=True, size=10)
            first += span
        self.ws.row_dimensions[self.row].height = max(32, max(len(wrapped_lines(header, width)) for header, (_, _, width) in zip(headers, columns)) * 16 + 12)
        self.row += 1
        for index, values in enumerate(rows):
            wrapped = [wrapped_lines(value, width) for value, (_, _, width) in zip(values, columns)]
            for start in range(0, max(map(len, wrapped)), 12):
                height = 1
                for lines, (first, last, _) in zip(wrapped, columns):
                    chunk = lines[start:start + 12]
                    height = max(height, len(chunk))
                    self.cell(self.row, first, last, '\n'.join(chunk), bg='F6F8FB' if index % 2 == 0 else 'FFFFFF', size=10)
                self.ws.row_dimensions[self.row].height = max(32, height * 16 + 12)
                self.row += 1
        if not rows:
            self.text('등록된 내역이 없습니다.')

    def finish(self):
        self.ws.print_area = f'A1:{get_column_letter(COLUMNS)}{self.row - 1}'


def request_sections(doc):
    """The request tab: system, background, type-specific inputs, and notes."""
    yield '대상 시스템', doc.get('related_system')
    yield '요청 배경', doc.get('background')
    request_type = doc.get('request_type')
    detail = doc.get('type_detail') or {}
    editor_key = EDITOR_FIELDS.get(request_type)
    keys = list(dict.fromkeys([*REQUEST_FIELDS.get(request_type, []), *detail]))
    for key in keys:
        value = (doc.get('description') or detail.get(key)) if key == editor_key else detail.get(key)
        value = DETAIL_OPTIONS.get(request_type, {}).get(key, {}).get(str(value), value)
        if value and key in {'occurredAt', 'applyDatetime', 'workDatetime'}:
            value = date_text(value, time=True)
        elif value and key in {'dataPeriodFrom', 'dataPeriodTo', 'retentionDate', 'permissionExpiry', 'actionDeadline', 'expiryDate'}:
            value = date_text(value)
        label = '검증 방법' if request_type == 'CONFIG_CHANGE' and key == 'verificationMethod' else DETAIL_LABELS.get(key, key)
        yield label, value
    if not editor_key:
        yield '요청 상세 내용', doc.get('description')
    yield '비고', doc.get('note')


def build_workbook(doc, sections, attachments):
    wb = Workbook()
    wb.remove(wb.active)
    sr_no, title = doc.get('sr_no') or str(doc['_id']), doc.get('title') or 'SR 요청'
    sheet = ReportSheet(wb, '요청 내용', sr_no, title, REQUEST_TYPE_LABEL.get(doc.get('request_type'), 'SR 요청'))
    for label, value in sections:
        if value is None or value == '' or value == [] or value == {}:
            continue
        if isinstance(value, list) and all(isinstance(row, dict) for row in value):
            keys = list(dict.fromkeys(k for row in value for k in row))
            sheet.section(label)
            for start in range(0, len(keys), 6):
                subset = keys[start:start + 6]
                sheet.table([DETAIL_LABELS.get(k, k) for k in subset],
                            [[plain_text(row.get(k, '')) for k in subset] for row in value])
            # Images in table cells are kept beside their table on the same sheet.
            for row in value:
                for cell in row.values():
                    for image in rich_root(cell).xpath('.//img[@src]'):
                        sheet.picture(attachments.add(label, url=image.get('src')))
        else:
            sheet.detail(label, value, attachments)

    extra = [item for item in attachments.items if item.number not in sheet.rendered_images]
    if extra:
        sheet.section('첨부파일')
    for item in extra:
        sheet.band(item.name, fg=BLUE, size=10)
        if item.image:
            sheet.picture(item)
        elif item.warning:
            sheet.text(item.warning)
        if item.bundled:
            cell = sheet.cell(sheet.row, 1, COLUMNS, item.archive_name, fg=BLUE, size=10)
            cell.hyperlink = item.archive_name
            sheet.ws.row_dimensions[sheet.row].height = max(30, len(wrapped_lines(item.archive_name)) * 17 + 12)
            sheet.row += 1
    sheet.finish()
    return wb


def export_detail(doc):
    sections = list(request_sections(doc))
    attachments = Attachments()
    for item in doc.get('attachments') or []:
        attachments.add('요청 첨부파일', item)
    for label, value in sections:
        attachments.scan(value, label)
    name = 'SR상세_' + safe_filename(doc.get('sr_no') or str(doc['_id']))
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
    workbook = None
    try:
        # Images are embedded in Excel. A ZIP is only needed for other files.
        if any(not item.image for item in attachments.items):
            with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
                attachments.prepare(archive)
                workbook = build_workbook(doc, sections, attachments)
                with SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b') as excel:
                    workbook.save(excel)
                    excel.seek(0)
                    with archive.open(name + '.xlsx', 'w', force_zip64=True) as member:
                        while chunk := excel.read(1024 * 1024):
                            member.write(chunk)
            media_type, extension = 'application/zip', '.zip'
        else:
            attachments.prepare()
            workbook = build_workbook(doc, sections, attachments)
            workbook.save(output)
            media_type, extension = XLSX_TYPE, '.xlsx'
        output.seek(0)
        return output, name + extension, media_type, sum(bool(item.warning) for item in attachments.items)
    except Exception:
        output.close()
        raise
    finally:
        if workbook:
            workbook.close()


def stream_export(output):
    try:
        while chunk := output.read(1024 * 1024):
            yield chunk
    finally:
        output.close()
