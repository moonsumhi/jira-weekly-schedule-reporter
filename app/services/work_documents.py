"""Work documents: ordered Markdown import, durable snapshots and HWPX export."""
import base64
import io
import logging
import re
import subprocess
import tempfile
import uuid
from html import escape as escape_html
from pathlib import Path
from urllib.parse import unquote, urlparse
from zipfile import ZipFile

from lxml import html
from markdown_it import MarkdownIt
from PIL import Image, UnidentifiedImageError

UPLOAD_ROOT = Path('/app/uploads')
logger = logging.getLogger(__name__)
DOCUMENT_SECTION = '문서 본문'


class DocumentImportError(ValueError):
    """An import problem with a message safe to display to the user."""


def _image_import_location(node) -> str:
    """Identify the surrounding row and column without exposing file paths."""
    cell = next((parent for parent in node.iterancestors() if parent.tag in ('td', 'th')), None)
    if cell is None:
        return ''
    row = cell.getparent()
    cells = row.xpath('./td|./th')
    if cell not in cells:
        return ''
    column = cells.index(cell)
    labels = [' '.join(item.itertext()).strip() for item in cells[:column]]
    labels = [re.sub(r'\s+', ' ', label) for label in labels if label and not label.isdecimal()]
    row_label = ' / '.join(labels[:2])[:100]
    table = next(cell.iterancestors('table'), None)
    rows = table.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr') if table is not None else []
    headers = rows[0].xpath('./td|./th') if rows and rows[0] is not row else []
    # Avoid assigning a misleading column name for merged or multi-level headers.
    column_label = ''
    if len(headers) == len(cells) and all(
        item.get('colspan', '1') == '1' and item.get('rowspan', '1') == '1'
        for item in headers + cells
    ):
        column_label = re.sub(r'\s+', ' ', ' '.join(headers[column].itertext())).strip()[:60]
    return ' → '.join(part for part in (row_label, column_label) if part)


def escape(text: str) -> str:
    return re.sub(r'([\\`*_{}\[\]<>#+!|])', r'\\\1', text)


def store_image(raw: bytes) -> str:
    image = Image.open(io.BytesIO(raw))
    image.load()
    image.thumbnail((2400, 2400))
    output = io.BytesIO()
    image.save(output, format='PNG')
    directory = UPLOAD_ROOT / 'work_documents' / 'images'
    directory.mkdir(parents=True, exist_ok=True)
    name = uuid.uuid4().hex + '.png'
    (directory / name).write_bytes(output.getvalue())
    return '/api/uploads/work_documents/images/' + name


def html_markdown(source: str, image_reader) -> str:
    """Walk in document order, including images between text inside table cells."""
    root = html.fragment_fromstring(source, create_parent='div')

    def walk(node):
        if not isinstance(node.tag, str):
            return ''
        tag = node.tag.lower() if isinstance(node.tag, str) else ''
        if tag in ('script', 'style', 'head'):
            return ''
        if tag == 'img':
            src = node.get('src', '')
            try:
                image_url = image_reader(src)
            except (FileNotFoundError, UnidentifiedImageError) as exc:
                location = _image_import_location(node)
                context = f' 위치: {location}.' if location else ''
                raise DocumentImportError(
                    f'문서의 이미지를 불러올 수 없어 가져오기를 중단했습니다.{context} '
                    '한글에서 해당 위치의 이미지 개체를 삭제하거나 다시 삽입한 뒤 저장해 주세요.'
                ) from exc
            return f'![사진](<{image_url}>)'
        if tag == 'br':
            return '\n'
        if tag == 'a':
            label = ''.join(node.itertext())
            href = node.get('href', '')
            if urlparse(href).scheme in ('http', 'https', 'mailto'):
                return f'[{escape(label)}](<{href.replace(">", "%3E").replace("<", "%3C")}>)'
            return escape(label)
        if tag == 'table':
            source_rows = node.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
            if not source_rows:
                return ''
            cell_rows = [row.xpath('./td|./th') for row in source_rows]
            def label(cell):
                return re.sub(r'\s+', '', ''.join(cell.itertext()))
            # Omit only the cover approval table, keeping the document title.
            # Worker/reviewer tables in the body remain part of the document.
            cover_labels = {label(cell) for cell in cell_rows[0]}
            cover_title = next((value for value in cover_labels if re.fullmatch(
                r'작업(?:계획|결과)서(?:\(서비스(?:외)?\))?', value)), None)
            if (cover_title and {'작업자', '담당자', '데이터운영팀담당자', '데이터운영팀장'} <= cover_labels
                    and not node.xpath('preceding::table|ancestor::table')):
                return '\n\n# ' + cover_title + '\n\n'
            # These job forms use merged cells to arrange label/value pairs,
            # not an eleven-column dataset. Keep each value with its label.
            if cell_rows[0] and label(cell_rows[0][0]) == '작업명':
                labels = {'작업명', '작업일시', '서비스명', '회사명/성함/직책', '중요도', '목적', '구분', '서비스영향도'}
                parts = []
                for cells in cell_rows:
                    current, values = None, []
                    for cell in cells:
                        if label(cell) in labels:
                            if current is not None:
                                parts.extend([f'### {current}', (' ' if current == '구분' else '\n\n').join(values)])
                            current, values = ''.join(cell.itertext()).strip(), []
                        else:
                            values.append(walk(cell).strip())
                    if current is not None:
                        parts.extend([f'### {current}', (' ' if current == '구분' else '\n\n').join(values)])
                return '\n\n' + '\n\n'.join(parts) + '\n\n'
            # A nested test table cannot live inside a Markdown pipe-table cell.
            # Lift the containing record into a section so its text, pictures,
            # and inner tables remain in their original order.
            if node.xpath('.//table'):
                header = cell_rows[0]
                has_header = all(not cell.xpath('.//table|.//img') and len(label(cell)) < 40 for cell in header)
                headers = [walk(cell).strip() for cell in header] if has_header else []
                parts = []
                for index, cells in enumerate(cell_rows[1:] if has_header else cell_rows, 1):
                    parts.append(f'### {index}번째 항목')
                    compact = []
                    def flush():
                        if compact:
                            parts.append('| ' + ' | '.join(name for name, _ in compact) + ' |\n| ' + ' | '.join('---' for _ in compact) + ' |\n| ' + ' | '.join(value for _, value in compact) + ' |')
                            compact.clear()
                    for c, cell in enumerate(cells):
                        name = headers[c] if c < len(headers) else f'{c + 1}번째 내용'
                        value = walk(cell).strip()
                        if not cell.xpath('.//table|.//img') and len(value) < 160:
                            compact.append((name.replace('\n', ' '), value.replace('\n', '<br>')))
                        else:
                            flush()
                            parts.extend([f'#### {name}', value])
                    flush()
                return '\n\n' + '\n\n'.join(parts) + '\n\n'
            grid = {}
            owners = {}
            def span(cell, name):
                try:
                    return max(1, min(100, int(cell.get(name, '1'))))
                except ValueError:
                    return 1
            for r, row in enumerate(source_rows):
                c = 0
                for cell in row.xpath('./td|./th'):
                    colspan, rowspan = span(cell, 'colspan'), span(cell, 'rowspan')
                    while any((r, c + offset) in grid for offset in range(colspan)):
                        c += 1
                    # Paragraph boundaries must not become four line breaks per cell.
                    value = re.sub(r'\n[ \t]*\n+', '\n', walk(cell).strip()).replace('\n', '<br>')
                    value = re.sub(r'(?<!\\)\|', r'\\|', value)
                    for dr in range(min(rowspan, len(source_rows) - r)):
                        for dc in range(colspan):
                            grid[r + dr, c + dc] = value if dr == dc == 0 else ''
                            owners[r + dr, c + dc] = value
                    c += colspan
            if not grid:
                return ''
            width = max(c for _, c in grid) + 1
            rows = [[grid.get((r, c), '') for c in range(width)] for r in range(len(source_rows))]
            # Collapse the two-level headings used by target, time, and
            # before/after tables. The second row is a heading, not data.
            second_labels = {label(cell) for cell in cell_rows[1]} if len(cell_rows) > 1 else set()
            is_grouped_header = ({'IP', 'HOSTNAME'} <= second_labels or {'시작', '종료'} <= second_labels
                                 or {'작업전', '작업후'} <= second_labels)
            if is_grouped_header and len(rows) > 1:
                rows = [[owners.get((1, c), '') or owners.get((0, c), '') for c in range(width)]] + rows[2:]
                if {'시작', '종료'} <= second_labels:
                    for c, value in enumerate(rows[0]):
                        parent = owners.get((0, c), '')
                        if parent and parent != value:
                            rows[0][c] = parent + ' / ' + value
                if not rows[0][0] and {'IP', 'HOSTNAME'} <= second_labels:
                    rows[0][0] = 'No.'
            title = ''
            for c, value in enumerate(rows[0]):
                normalized = re.sub(r'\s+|<br>', '', value)
                if width > 1 and re.fullmatch(r'작업(?:계획|결과)서(?:\(서비스(?:외)?\))?', normalized):
                    title = '# ' + normalized + '\n\n'
                    rows = [row[:c] + row[c + 1:] for row in rows]
                    width -= 1
                    break
            result = ['| ' + ' | '.join(row) + ' |' for row in rows]
            result.insert(1, '| ' + ' | '.join(['---'] * width) + ' |')
            return '\n\n' + title + '\n'.join(result) + '\n\n'
        def text(value):
            # Ignore pretty-printed XHTML indentation between block elements.
            return '' if value and not value.strip() and '\n' in value else escape(value or '')
        value = text(node.text)
        for child in node:
            value += walk(child) + text(child.tail)
        if tag in ('strong', 'b'):
            return '**' + value + '**'
        if tag in ('em', 'i'):
            return '*' + value + '*'
        if re.fullmatch(r'h[1-6]', tag):
            return '\n\n' + '#' * int(tag[1]) + ' ' + value.strip() + '\n\n'
        if tag == 'li':
            return '\n- ' + value.strip() + '\n'
        if tag in ('p', 'div', 'ul', 'ol'):
            section = re.fullmatch(r'\[([^\[\]\n]+)\]', ''.join(node.itertext()).strip()) if tag == 'p' else None
            if section and not node.xpath('.//table|.//img'):
                return '\n\n## ' + escape(section.group(1)) + '\n\n'
            return '\n\n' + value.strip() + '\n\n'
        return value

    return re.sub(r'\n{3,}', '\n\n', walk(root)).strip() + '\n'


def import_document(content: bytes, filename: str) -> tuple[str, list[str]]:
    suffix = Path(filename).suffix.lower()
    warnings = ['원본의 글꼴·페이지 배치·병합 표는 달라질 수 있습니다. 저장 전에 내용을 확인해 주세요.']
    if suffix == '.hwpx':
        from lxml import etree
        with ZipFile(io.BytesIO(content)) as archive:
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            manifest = etree.fromstring(archive.read('Contents/content.hpf'), parser)
            paths = {}
            for item in manifest.iter():
                href = item.get('href', '')
                if item.get('id') and 'BinData/' in href:
                    paths[item.get('id')] = href[href.index('BinData/'):]
            def convert(node):
                tag = etree.QName(node).localname
                if tag == 'pic':
                    for image in node.iter():
                        key = image.get('binaryItemIDRef')
                        if key and key in paths:
                            return '<img src="' + store_image(archive.read(paths[key])) + '"/>'
                    raise ValueError('한글 문서의 사진을 읽지 못했습니다.')
                if tag == 't':
                    return escape_html(node.text or '') + ''.join(convert(child) for child in node)
                if tag == 'lineBreak':
                    return '<br>'
                if tag in ('secPr', 'ctrl', 'linesegarray'):
                    return ''
                text = ''.join(convert(child) for child in node)
                markup = {'p': 'p', 'tbl': 'table', 'tr': 'tr', 'tc': 'td'}.get(tag)
                if tag == 'tc':
                    spans = node.xpath('./*[local-name()="cellSpan"]')
                    attributes = ''
                    if spans:
                        for xml_name, html_name in (('colSpan', 'colspan'), ('rowSpan', 'rowspan')):
                            value = spans[0].get(xml_name, '1')
                            if value.isdigit():
                                attributes += f' {html_name}="{value}"'
                    return f'<td{attributes}>{text}</td>'
                return f'<{markup}>{text}</{markup}>' if markup else text
            names = sorted((name for name in archive.namelist() if re.fullmatch(r'Contents/section\d+\.xml', name)),
                           key=lambda name: int(re.search(r'\d+', name).group()))
            source = ''.join(convert(etree.fromstring(archive.read(name), parser)) for name in names)
            return html_markdown(source, lambda src: src), warnings
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        source = root / ('source' + suffix)
        source.write_bytes(content)
        if suffix == '.doc':
            subprocess.run(['libreoffice', '-env:UserInstallation=file://' + str(root / 'profile'),
                            '--headless', '--convert-to', 'docx', '--outdir', str(root), str(source)],
                           check=True, capture_output=True, timeout=90)
            source = root / 'source.docx'
            suffix = '.docx'
        if suffix == '.docx':
            import mammoth
            def convert(image):
                with image.open() as stream:
                    return {'src': store_image(stream.read())}
            with source.open('rb') as stream:
                result = mammoth.convert_to_html(stream, convert_image=mammoth.images.img_element(convert))
            warnings.extend(message.message for message in result.messages)
            return html_markdown(result.value, lambda src: src), warnings
        if suffix == '.hwp':
            output = root / 'converted'
            subprocess.run(['hwp5html', '--output', str(output), str(source)],
                           check=True, capture_output=True, timeout=90)
            index = output / 'index.xhtml'
            if not index.exists():
                raise ValueError('한글 문서 HTML 변환 결과가 없습니다.')
            def read_image(src):
                if src.startswith('data:'):
                    return store_image(base64.b64decode(src.split(',', 1)[1]))
                path = (output / unquote(src)).resolve()
                if not path.is_relative_to(output.resolve()):
                    raise ValueError('잘못된 이미지 경로입니다.')
                return store_image(path.read_bytes())
            return html_markdown(index.read_text(encoding='utf-8'), read_image), warnings
    raise ValueError('HWP, HWPX, DOC, DOCX 파일을 선택해 주세요.')


def markdown_from_data(data: dict) -> str | None:
    rows = data.get(DOCUMENT_SECTION)
    if isinstance(rows, list) and len(rows) == 1 and isinstance(rows[0], dict):
        row = rows[0]
        if row.get('내용__format') == 'markdown' and isinstance(row.get('내용'), str):
            return row['내용']
    return None


def save_markdown_snapshot(data: dict) -> dict:
    markdown = markdown_from_data(data)
    if markdown is None:
        return {}
    directory = UPLOAD_ROOT / 'work_documents' / 'markdown'
    directory.mkdir(parents=True, exist_ok=True)
    name = uuid.uuid4().hex + '.md'
    (directory / name).write_text(markdown, encoding='utf-8')
    return {'document_format': 'markdown', 'markdown_file': str(directory / name)}


def local_image(src: str) -> bytes:
    parsed = urlparse(src)
    # Never fetch user-controlled URLs: only embed files from our upload directory.
    path = unquote(parsed.path)
    if not path.startswith('/api/uploads/'):
        raise ValueError('서버에 업로드된 사진만 한글 파일에 포함할 수 있습니다.')
    target = (UPLOAD_ROOT / path[len('/api/uploads/'):]).resolve()
    if not target.is_relative_to(UPLOAD_ROOT.resolve()) or not target.is_file():
        raise ValueError('사진 파일을 찾을 수 없습니다. 상세 화면에서 사진을 다시 첨부해 주세요.')
    return target.read_bytes()


def _table_widths_for_export(node, rows, cols, nested=False):
    """Resolve explicit table widths and compact selected nested columns.

    Work-form tables carry their widths in ``data-column-widths``. Markdown
    tables embedded inside a cell do not, so use their header labels to keep
    short ``구분``/``테스트 항목`` columns compact and give the saved width to
    ``비고`` when it is present.
    """
    raw_widths = node.get('data-column-widths', '')
    try:
        parsed_widths = [float(value.strip()) for value in raw_widths.split(',') if value.strip()]
        if len(parsed_widths) == cols and any(parsed_widths):
            return parsed_widths
    except (TypeError, ValueError):
        logger.debug('Ignoring invalid table width hint: %r', raw_widths)

    widths = [100.0 / cols] * cols
    if not nested or not rows:
        return widths

    header = rows[0].xpath('./td|./th')
    labels = [re.sub(r'\s+', '', ''.join(cell.itertext())) for cell in header]
    # Keep these close to the width of their short headers. The saved width
    # is handed to 비고, which is usually the long free-text column.
    targets = {'구분': 10.0, '테스트항목': 19.0}
    compacted = []
    for label, target_width in targets.items():
        try:
            index = labels.index(label)
        except ValueError:
            continue
        if widths[index] > target_width:
            compacted.append((index, widths[index] - target_width))
            widths[index] = target_width
    if not compacted:
        return widths

    receiver = next((index for index, label in enumerate(labels) if label == '비고'), None)
    if receiver is None:
        # Keep the table at 100% even when a nested table has no 비고 column.
        receiver = max((index for index in range(cols)
                        if index not in {item[0] for item in compacted}),
                       key=lambda index: widths[index], default=None)
    if receiver is not None:
        widths[receiver] += sum(amount for _, amount in compacted)
    return widths


def export_hwpx(markdown: str) -> bytes:
    from hwpx import HwpxDocument
    document = HwpxDocument.new()
    # Match the document layout used by the work-document editor.  The HWPX
    # skeleton otherwise keeps its wider default margins (about 3 cm), which
    # leaves too little room for the exported tables.  Set all page margins to
    # 0 mm vertically and 20 mm horizontally so the table width below
    # (170 mm on A4) fits the usable page area without wasting vertical space.
    document.page.setup(
        margins_mm={
            'left': 20,
            'right': 20,
            'top': 0,
            'bottom': 0,
        }
    )
    rendered = MarkdownIt('commonmark', {'html': True}).enable('table').render(markdown)
    root = html.fragment_fromstring(rendered, create_parent='div')
    initial_paragraph = document.paragraphs[0] if document.paragraphs else None
    # Keep the document title and all table headings centered while leaving
    # table values in their default left alignment.  Reuse the same paragraph
    # properties for every heading so HWPX and the generated HWP stay
    # consistent across sections and nested tables.
    center_para_pr_ref = None
    if initial_paragraph is not None:
        document.styles.apply_paragraph_format(paragraph_index=0, alignment='CENTER')
        center_para_pr_ref = document.paragraphs[0].para_pr_id_ref
    first_paragraph = True

    def add_document_paragraph(text='', **kwargs):
        """Reuse the skeleton paragraph so exported titles start at the top."""
        nonlocal first_paragraph
        if first_paragraph and initial_paragraph is not None:
            first_paragraph = False
            return initial_paragraph
        first_paragraph = False
        return document.add_paragraph(text, **kwargs)

    def emit(node, add_paragraph, width_mm=170, font_size=None):
        tag = node.tag.lower() if isinstance(node.tag, str) else ''
        if tag in ('script', 'style'):
            return
        if tag == 'table':
            rows = node.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
            if not rows:
                return
            def span(cell, name='colspan'):
                try:
                    return max(1, int(cell.get(name, '1')))
                except (TypeError, ValueError):
                    return 1

            cols = max(sum(span(cell) for cell in row.xpath('./td|./th')) for row in rows)
            # The HWPX library's default row height is 3600 HWP units, which
            # leaves a large blank band above and below one-line values. Use a
            # compact one-line baseline; rows with longer content can still
            # expand when Hancom lays the table out.
            table = add_paragraph('').add_table(
                rows=len(rows), cols=cols,
                width=round(width_mm * 7200 / 25.4),
                height=2200 * len(rows),
            )
            # HWP's default table object is marked "글자처럼 취급"
            # (treatAsChar=1), which makes a long table behave like one very
            # wide character. Use a flowing table anchor so Hancom can keep
            # the table inside the page/column layout and split it naturally
            # across pages when the content grows.
            table_pos = table.element.find(
                '{http://www.hancom.co.kr/hwpml/2011/paragraph}pos'
            )
            if table_pos is not None:
                table_pos.set('treatAsChar', '0')
                table_pos.set('flowWithText', '1')
                table_pos.set('allowOverlap', '0')
                table_pos.set('vertRelTo', 'PARA')
                table_pos.set('horzRelTo', 'COLUMN')
                table_pos.set('vertAlign', 'TOP')
                table_pos.set('horzAlign', 'LEFT')
                table_pos.set('vertOffset', '0')
                table_pos.set('horzOffset', '0')
            widths = _table_widths_for_export(node, rows, cols, nested=font_size is not None)
            table.set_column_widths(widths)
            total_weight = sum(widths)

            def cell_width(start, span):
                weight = sum(widths[start:min(start + span, cols)])
                return max(10, width_mm * weight / total_weight - 4)

            for r, row in enumerate(rows):
                c = 0
                for cell in row.xpath('./td|./th'):
                    colspan = span(cell)
                    if colspan > 1:
                        table.merge_cells(r, c, r, min(cols - 1, c + colspan - 1))
                    target = table.cell(r, c)
                    role = cell.get('data-role', '')
                    if role == 'section-heading':
                        table.set_cell_shading(r, c, 'D9D9D9')
                    elif cell.tag.lower() == 'th':
                        table.set_cell_shading(r, c, 'F2F2F2')
                    # A new HWPX table cell contains one empty paragraph by
                    # default. Remove it before emitting the first header or
                    # value so the content does not start one line too low.
                    has_content = bool(''.join(cell.itertext()).strip() or cell.xpath('.//img|.//table'))
                    if has_content:
                        for paragraph in target.paragraphs:
                            if not paragraph.text and not paragraph.tables:
                                parent = paragraph.element.getparent()
                                if parent is not None:
                                    parent.remove(paragraph.element)
                                break
                    # Tables nested inside a detailed-work cell use a compact
                    # 9pt header / 8pt data scale; top-level report tables
                    # remain 10pt / 9pt.
                    nested = font_size is not None
                    emit(cell, target.add_paragraph, cell_width(c, colspan),
                         font_size=(9 if nested else 10)
                         if cell.tag.lower() == 'th' else (8 if nested else 9))
                    if cell.tag.lower() == 'th' and center_para_pr_ref is not None:
                        for paragraph in target.paragraphs:
                            paragraph.para_pr_id_ref = center_para_pr_ref
                    c += colspan
            return
        paragraph = None
        def text(value, bold=False, italic=False):
            nonlocal paragraph
            if value:
                if paragraph is None:
                    paragraph = add_paragraph('')
                    if center_para_pr_ref is not None and re.fullmatch(r'h[1-6]', tag):
                        paragraph.para_pr_id_ref = center_para_pr_ref
                heading = int(tag[1]) if re.fullmatch(r'h[1-6]', tag) else 0
                size = font_size if font_size is not None else (max(12, 22 - heading * 2) if heading else 10)
                style = document.styles.ensure_run(bold=bold or bool(heading), italic=italic,
                                                   size=size)
                paragraph.add_run(value, char_pr_id_ref=style)
        def inline(element, bold=False, italic=False):
            nonlocal paragraph
            if element.tag == 'img':
                raw = local_image(element.get('src', ''))
                picture = Image.open(io.BytesIO(raw))
                picture.load()
                buf = io.BytesIO()
                picture.save(buf, format='PNG')
                # The public document API embeds BinData, then move its picture run
                # into the destination paragraph (also works inside table cells).
                image_width = min(width_mm, 200 * picture.width / picture.height)
                holder = document.add_picture(buf.getvalue(), 'png', width_mm=image_width,
                                              height_mm=image_width * picture.height / picture.width)
                target = add_paragraph('')
                run = holder.element.getparent()
                source_p = run.getparent()
                source_p.remove(run)
                target.element.append(run)
                source_p.getparent().remove(source_p)
                paragraph = None
                return
            if element.tag == 'br':
                paragraph = None
                return
            if element.tag in ('script', 'style'):
                return
            bold = bold or element.tag in ('b', 'strong', 'th')
            italic = italic or element.tag in ('i', 'em')
            if element.tag == 'li':
                text('• ')
            text(element.text, bold, italic)
            for child in element:
                inline(child, bold, italic)
                text(child.tail, bold, italic)
        if tag in ('div', 'ul', 'ol', 'blockquote') or (tag in ('td', 'th') and any(child.tag in ('p', 'table', 'div', 'ul', 'ol') for child in node)):
            if node.text and node.text.strip():
                add_paragraph(node.text)
            for child in node:
                emit(child, add_paragraph, width_mm, font_size=font_size)
            return
        inline(node)

    for child in root:
        emit(child, add_document_paragraph)
    stream = io.BytesIO()
    document.save_to_stream(stream)
    return stream.getvalue()


def export_hwp(markdown: str) -> bytes:
    """Produce an actual HWP 5 binary for Hancom Office 2010."""
    with tempfile.TemporaryDirectory() as temporary:
        source = Path(temporary) / 'document.hwpx'
        target = Path(temporary) / 'document.hwp'
        source.write_bytes(export_hwpx(markdown))
        # The HWPX already contains the report styling and explicit column
        # widths.  Running hwp edit --style-tables here would recalculate
        # widths from the current values and shrink empty columns such as
        # "비고" to one or two characters.
        subprocess.run(['hwp', 'convert', str(source), '-o', str(target)],
                       check=True, capture_output=True, timeout=90)
        output = target.read_bytes()
        if not output.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
            raise ValueError('올바른 HWP 파일을 생성하지 못했습니다.')
        return output


def export_docx(markdown: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Noto Sans CJK KR', Pt(10)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK KR')
    root = html.fragment_fromstring(MarkdownIt('commonmark', {'html': True}).enable('table').render(markdown), create_parent='div')

    def emit(node, container, width=170, font_size=None):
        tag = node.tag.lower() if isinstance(node.tag, str) else ''
        if tag in ('script', 'style'):
            return
        if tag == 'table':
            rows = node.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
            if not rows:
                return
            def span(cell, name='colspan'):
                try:
                    return max(1, int(cell.get(name, '1')))
                except (TypeError, ValueError):
                    return 1

            cols = max(sum(span(cell) for cell in row.xpath('./td|./th')) for row in rows)
            table = container.add_table(rows=len(rows), cols=cols)
            table.style = 'Table Grid'
            table.autofit = True
            from docx.oxml import OxmlElement
            widths = _table_widths_for_export(node, rows, cols, nested=font_size is not None)
            total_weight = sum(widths)

            def cell_width(start, span):
                weight = sum(widths[start:min(start + span, cols)])
                return max(10, width * weight / total_weight - 4)

            for r, row in enumerate(rows):
                c = 0
                for cell in row.xpath('./td|./th'):
                    colspan = span(cell)
                    target = table.cell(r, c)
                    if colspan > 1:
                        target = target.merge(table.cell(r, min(cols - 1, c + colspan - 1)))
                    role = cell.get('data-role', '')
                    if role == 'section-heading' or cell.tag.lower() == 'th':
                        shading = OxmlElement('w:shd')
                        shading.set(qn('w:fill'), 'D9D9D9' if role == 'section-heading' else 'F2F2F2')
                        target._tc.get_or_add_tcPr().append(shading)
                    has_content = bool(''.join(cell.itertext()).strip() or cell.xpath('.//img|.//table'))
                    if has_content:
                        paragraphs = target.paragraphs
                        if paragraphs and not paragraphs[0].text and not paragraphs[0]._element.xpath('.//w:tbl'):
                            paragraphs[0]._element.getparent().remove(paragraphs[0]._element)
                    nested = font_size is not None
                    emit(cell, target, cell_width(c, colspan),
                         font_size=(9 if nested else 10)
                         if cell.tag.lower() == 'th' else (8 if nested else 9))
                    if cell.tag.lower() == 'th':
                        for paragraph in target.paragraphs:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    c += colspan
            return
        if tag in ('div', 'ul', 'ol', 'blockquote') or (tag in ('td', 'th') and any(child.tag in ('p', 'table', 'div', 'ul', 'ol') for child in node)):
            if node.text and node.text.strip():
                container.add_paragraph(node.text)
            for child in node:
                emit(child, container, width, font_size=font_size)
            return
        heading = int(tag[1]) if re.fullmatch(r'h[1-6]', tag) else 0
        paragraph = container.add_paragraph(style=f'Heading {heading}' if heading else None)
        if heading:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        def inline(element, bold=False, italic=False):
            if element.tag in ('script', 'style'):
                return
            if element.tag == 'img':
                raw = local_image(element.get('src', ''))
                with Image.open(io.BytesIO(raw)) as photo:
                    photo.load()
                    stream = io.BytesIO()
                    photo.save(stream, format='PNG')
                    picture_width = min(width, 200 * photo.width / photo.height)
                paragraph.add_run().add_picture(io.BytesIO(stream.getvalue()), width=Mm(picture_width))
                return
            if element.tag == 'br':
                paragraph.add_run().add_break()
                return
            bold = bold or element.tag in ('strong', 'b', 'th')
            italic = italic or element.tag in ('em', 'i')
            def text(value):
                if value:
                    run = paragraph.add_run(value)
                    run.bold, run.italic = bold, italic
                    if font_size is not None:
                        run.font.size = Pt(font_size)
            if element.tag == 'li':
                text('• ')
            text(element.text)
            for child in element:
                inline(child, bold, italic)
                text(child.tail)
        inline(node)

    for child in root:
        emit(child, document)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
