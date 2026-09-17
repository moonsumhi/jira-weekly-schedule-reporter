import io
import tempfile
import subprocess
import unittest
from pathlib import Path
from unittest.mock import patch
from unittest.mock import AsyncMock
from types import SimpleNamespace
from zipfile import ZipFile

from PIL import Image
from lxml import etree, html
from app.services import work_documents as documents


class WorkDocumentTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        patcher = patch.object(documents, 'UPLOAD_ROOT', Path(self.directory.name))
        patcher.start()
        self.addCleanup(patcher.stop)

    def photo(self):
        stream = io.BytesIO()
        Image.new('RGB', (120, 80), 'blue').save(stream, 'PNG')
        return stream.getvalue()

    def test_html_cell_order(self):
        result = documents.html_markdown('<table><tr><td>앞<img src="photo.png"/>뒤</td></tr></table>', lambda _: '/api/uploads/a.png')
        self.assertLess(result.index('앞'), result.index('![사진]'))
        self.assertLess(result.index('![사진]'), result.index('뒤'))

    def test_merged_table_columns_and_paragraph_spacing(self):
        source = '''<?xml version="1.0" encoding="utf-8"?>
        <table><tr><td rowspan="2"><p>대상</p></td><td colspan="2"><p>서버 정보</p></td></tr>
        <tr><td><p>IP</p><p>주소</p></td><td>HOSTNAME</td></tr>
        <tr><td>서버1</td><td>10.0.0.1</td><td>host1</td></tr></table>'''
        markdown = documents.html_markdown(source, lambda src: src)
        self.assertNotIn('encoding=', markdown)
        self.assertIn('| 대상 | 서버 정보 |  |', markdown)
        self.assertIn('|  | IP<br>주소 | HOSTNAME |', markdown)
        self.assertNotIn('<br><br>', markdown)
        from markdown_it import MarkdownIt
        rendered = html.fromstring(MarkdownIt().enable('table').render(markdown))
        self.assertEqual([len(row) for row in rendered.xpath('.//tr')], [3, 3, 3])

    def test_nested_table_does_not_create_extra_columns(self):
        markdown = documents.html_markdown('<table><tr><td>항목</td><td>내용</td></tr><tr><td>서버</td><td><table><tr><td>IP</td><td>주소</td></tr></table></td></tr></table>', lambda src: src)
        from markdown_it import MarkdownIt
        rendered = html.fromstring(MarkdownIt().enable('table').render(markdown))
        self.assertEqual(len(rendered.xpath('.//tr')[-1]), 2)
        self.assertIn('주소', ''.join(rendered.itertext()))
        self.assertNotIn('\\|', markdown)

    def test_job_form_layout_and_grouped_headers(self):
        source = '<p>[작업 개요]</p><table><tr><td>작 업 명</td><td colspan="8">배포 작업</td><td>작업 일시</td><td>2026-08-20</td></tr><tr><td>목적</td><td colspan="10">개선</td></tr></table><p>[작업 대상]</p><table><tr><td rowspan="2">No.</td><td colspan="3">작업 대상</td><td rowspan="2">비고</td></tr><tr><td>작업 대상</td><td>IP</td><td>HOSTNAME</td></tr><tr><td>1</td><td>서버</td><td>주소</td><td>호스트</td><td>정상</td></tr></table>'
        markdown = documents.html_markdown(source, lambda src: src)
        self.assertIn('## 작업 개요', markdown)
        self.assertIn('### 작 업 명\n\n배포 작업', markdown)
        self.assertIn('| No. | 작업 대상 | IP | HOSTNAME | 비고 |', markdown)
        self.assertNotIn('|  |  |', markdown)
        time_table = '<table><tr><td colspan="2">작업 시간</td><td rowspan="2">내용</td></tr><tr><td>시작</td><td>종료</td></tr><tr><td>17:00</td><td>17:30</td><td>배포</td></tr></table>'
        self.assertIn('| 작업 시간 / 시작 | 작업 시간 / 종료 | 내용 |', documents.html_markdown(time_table, lambda src: src))

    def test_cover_approval_table_removed_but_body_reviewers_retained(self):
        source = '<table><tr><td>작업자</td><td>담당자</td><td rowspan="2">작 업 계 획 서<br>(서비스)</td><td>데이터운영팀 담당자</td><td>데이터운영<br>팀장</td></tr><tr><td>결재자1</td><td>결재자2</td><td></td><td></td></tr></table><p>[담당자]</p><table><tr><td>소속</td><td>성함</td></tr><tr><td>운영팀</td><td>본문 담당자</td></tr></table>'
        markdown = documents.html_markdown(source, lambda src: src)
        self.assertTrue(markdown.startswith('# 작업계획서(서비스)'))
        self.assertNotIn('결재자1', markdown)
        self.assertNotIn('데이터운영', markdown)
        self.assertIn('## 담당자', markdown)
        self.assertIn('본문 담당자', markdown)

    def test_nested_table_text_picture_order_is_preserved(self):
        source = '<table><tr><td>제목</td><td>세부 작업 내용</td></tr><tr><td>점검</td><td><p>작업 전</p><img src="a.png"/><table><tr><td>테스트 항목</td><td>결과</td></tr><tr><td>검증</td><td>통과</td></tr></table><p>작업 후</p></td></tr></table>'
        markdown = documents.html_markdown(source, lambda src: '/api/uploads/a.png')
        self.assertLess(markdown.index('작업 전'), markdown.index('![사진]'))
        self.assertLess(markdown.index('![사진]'), markdown.index('| 테스트 항목'))
        self.assertLess(markdown.index('| 검증'), markdown.index('작업 후'))
        self.assertNotIn('\\|', markdown)

    def test_docx_order_and_hwpx_embedded_picture(self):
        from docx import Document
        doc = Document()
        paragraph = doc.add_paragraph('BEFORE')
        paragraph.add_run().add_picture(io.BytesIO(self.photo()))
        paragraph.add_run('AFTER')
        stream = io.BytesIO()
        doc.save(stream)
        markdown, _ = documents.import_document(stream.getvalue(), 'test.docx')
        self.assertLess(markdown.index('BEFORE'), markdown.index('![사진]'))
        self.assertLess(markdown.index('![사진]'), markdown.index('AFTER'))
        self.assert_hwpx_order(markdown)

    def test_hwp_unreadable_image_reports_its_table_location(self):
        def convert(command, **kwargs):
            output = Path(command[command.index('--output') + 1])
            output.mkdir()
            if corrupt:
                (output / 'bindata').mkdir()
                (output / 'bindata' / 'BIN000A.bmp').write_bytes(b'invalid bitmap')
            (output / 'index.xhtml').write_text(
                '<p>[개발 내용]</p><table><tr><td>제목</td><td>세부 작업 내용</td></tr>'
                '<tr><td>마이그레이션</td><td><p>테스트 케이스</p>'
                '<table><tr><td>구분</td><td>테스트 항목</td><td>예상 결과</td><td>실제 결과</td></tr>'
                '<tr><td>기능</td><td>GitLab 웹/Git 접속</td><td>정상 접속</td>'
                '<td>Pass<img src="bindata/BIN000A.bmp"/></td></tr></table>'
                '<p>작업 후</p></td></tr></table>',
                encoding='utf-8',
            )

        for corrupt in (False, True):
            with self.subTest(corrupt=corrupt), patch.object(documents.subprocess, 'run', side_effect=convert):
                with self.assertRaises(documents.DocumentImportError) as error:
                    documents.import_document(b'hwp fixture', 'missing-image.hwp')
            message = str(error.exception)
            self.assertIn('기능 / GitLab 웹/Git 접속 → 실제 결과', message)
            self.assertIn('이미지 개체를 삭제하거나 다시 삽입', message)
            self.assertNotIn('BIN000A', message)
            self.assertNotIn('/tmp/', message)

    def test_hwp_valid_image_preserves_surrounding_content(self):
        def convert(command, **kwargs):
            output = Path(command[command.index('--output') + 1])
            (output / 'bindata').mkdir(parents=True)
            (output / 'bindata' / 'BIN0001.png').write_bytes(self.photo())
            (output / 'index.xhtml').write_text(
                '<table><tr><td>점검</td><td>결과</td></tr><tr><td>서비스</td>'
                '<td>작업 전<img src="bindata/BIN0001.png"/>작업 후</td></tr></table>', encoding='utf-8',
            )

        with patch.object(documents.subprocess, 'run', side_effect=convert):
            markdown, warnings = documents.import_document(b'hwp fixture', 'valid-image.hwp')
        self.assertLess(markdown.index('작업 전'), markdown.index('![사진]'))
        self.assertLess(markdown.index('![사진]'), markdown.index('작업 후'))
        self.assertIn('| 점검 | 결과 |', markdown)
        self.assertEqual(len(warnings), 1)
        self.assertEqual(len(list(Path(self.directory.name).rglob('*.png'))), 1)

    def test_hwp_missing_image_does_not_allow_path_traversal(self):
        def convert(command, **kwargs):
            output = Path(command[command.index('--output') + 1])
            output.mkdir()
            (output / 'index.xhtml').write_text('<p>본문<img src="%2e%2e/private.png"/></p>', encoding='utf-8')

        with patch.object(documents.subprocess, 'run', side_effect=convert):
            with self.assertRaisesRegex(ValueError, '잘못된 이미지 경로'):
                documents.import_document(b'hwp fixture', 'invalid-path.hwp')

    def assert_hwpx_order(self, markdown):
        output = documents.export_hwpx(markdown)
        from hwpx import HwpxDocument
        with HwpxDocument.open(output) as document:
            report = document.validate()
            self.assertTrue(report.ok, str(report.errors))
        reimported, _ = documents.import_document(output, 'roundtrip.hwpx')
        self.assertLess(reimported.index('BEFORE'), reimported.index('![사진]'))
        self.assertLess(reimported.index('![사진]'), reimported.index('AFTER'))
        with ZipFile(io.BytesIO(output)) as archive:
            self.assertIsNone(archive.testzip())
            root = etree.fromstring(archive.read('Contents/section0.xml'))
            ordered = []
            for node in root.iter():
                tag = etree.QName(node).localname
                if tag == 't' and node.text:
                    ordered.append(node.text)
                elif tag == 'pic':
                    ordered.append('PICTURE')
            text = ' '.join(ordered)
            self.assertLess(text.index('BEFORE'), text.index('PICTURE'))
            self.assertLess(text.index('PICTURE'), text.index('AFTER'))
            self.assertTrue(any(name.startswith('BinData/') for name in archive.namelist()))

    def test_table_picture_order(self):
        src = documents.store_image(self.photo())
        self.assert_hwpx_order(f'| 설명 |\n| --- |\n| BEFORE ![사진](<{src}>) AFTER |')

    def test_korean_utf8_is_preserved_in_hwpx(self):
        title = '작업계획서 한글 인코딩 확인'
        body = '작업 전 점검 → 작업 후 확인 · 담당자 홍길동'
        markdown = f'# {title}\n\n| 작업 내용 | 비고 |\n| --- | --- |\n| {body} | 정상 |'
        output = documents.export_hwpx(markdown)
        with ZipFile(io.BytesIO(output)) as archive:
            raw = archive.read('Contents/section0.xml')
            self.assertIn(b"encoding='UTF-8'", raw[:100])
            self.assertIn(title.encode('utf-8'), raw)
            self.assertIn(body.encode('utf-8'), raw)
            self.assertNotIn('\ufffd', raw.decode('utf-8'))
        reimported, _ = documents.import_document(output, 'encoding.hwpx')
        self.assertIn(title, reimported)
        self.assertIn(body, reimported)

    def test_hwpx_title_starts_in_the_first_paragraph(self):
        output = documents.export_hwpx('# 작업계획서\n\n본문')
        with ZipFile(io.BytesIO(output)) as archive:
            root = etree.fromstring(archive.read('Contents/section0.xml'))
        paragraphs = root.xpath('./*[local-name()="p"]')
        self.assertTrue(paragraphs)
        self.assertIn('작업계획서', ''.join(paragraphs[0].itertext()))

    def test_hwp_binary_preserves_korean_table_and_image(self):
        src = documents.store_image(self.photo())
        output = documents.export_hwp(f'# 한글 호환성\n\n| 작업 내용 | 비고 |\n| --- | --- |\n| 작업 전 ![사진](<{src}>) 작업 후 | 정상 |')
        self.assertTrue(output.startswith(bytes.fromhex('d0cf11e0a1b11ae1')))
        path = Path(self.directory.name) / 'compatibility.hwp'
        path.write_bytes(output)
        parsed = subprocess.run(['hwp5proc', 'xml', str(path)], check=True, capture_output=True, timeout=30)
        root = etree.fromstring(parsed.stdout)
        text = ''.join(root.itertext())
        self.assertIn('한글 호환성', text)
        self.assertIn('작업 전', text)
        self.assertIn('작업 후', text)
        self.assertIn('정상', text)
        self.assertTrue(root.xpath('.//TableControl'))
        self.assertTrue(root.xpath('.//PictureInfo'))

    def test_pdf_import_is_not_supported(self):
        with self.assertRaisesRegex(ValueError, 'HWP, HWPX, DOC, DOCX'):
            documents.import_document(b'%PDF-1.7', 'scan.pdf')
        from app.routers.form_entries import _ensure_work_document_import_format
        from fastapi import HTTPException
        with self.assertRaises(HTTPException) as error:
            _ensure_work_document_import_format('scan.pdf')
        self.assertEqual(error.exception.status_code, 415)

    def test_docx_export_preserves_korean_table_picture_order(self):
        src = documents.store_image(self.photo())
        output = documents.export_docx(f'# 작업계획서\n\n| 작업 내용 | 비고 |\n| --- | --- |\n| 작업 전 ![사진](<{src}>) 작업 후 | 정상 |')
        with ZipFile(io.BytesIO(output)) as archive:
            root = etree.fromstring(archive.read('word/document.xml'))
            ordered = [node.text if etree.QName(node).localname == 't' else 'PICTURE'
                       for node in root.iter() if etree.QName(node).localname in ('t', 'drawing')]
            text = ' '.join(value for value in ordered if value)
            self.assertIn('작업계획서', text)
            self.assertLess(text.index('작업 전'), text.index('PICTURE'))
            self.assertLess(text.index('PICTURE'), text.index('작업 후'))
            self.assertTrue(any(name.startswith('word/media/') for name in archive.namelist()))
            self.assertTrue(root.xpath('//*[local-name()="tbl"]'))
        for src in ('https://example.com/private.png', '/api/uploads/../../secret.png'):
            with self.assertRaises(ValueError):
                documents.export_docx(f'![사진]({src})')

    def test_markdown_snapshot(self):
        markdown = '앞\n\n![사진](/api/uploads/a.png)\n\n뒤'
        data = {'문서 본문': [{'내용': markdown, '내용__format': 'markdown'}]}
        saved = documents.save_markdown_snapshot(data)
        self.assertEqual(Path(saved['markdown_file']).read_text(encoding='utf-8'), markdown)

    def test_export_rejects_external_or_traversal_images(self):
        for src in ('http://localhost/secret.png', '/api/uploads/../../secret.png'):
            with self.assertRaises(ValueError):
                documents.local_image(src)


class WorkDocumentApiTests(unittest.IsolatedAsyncioTestCase):
    async def test_import_errors_explain_known_failures_without_exposing_internal_errors(self):
        from bson import ObjectId
        from fastapi import HTTPException, UploadFile
        from app.routers import form_entries
        collection = AsyncMock()
        collection.find_one.return_value = {'sections': []}
        user = SimpleNamespace(full_name='테스트', email='test@example.com')
        message = '문서의 이미지를 불러올 수 없습니다. 위치: 테스트 항목 → 실제 결과.'
        for known in (True, False):
            error = documents.DocumentImportError(message) if known else RuntimeError('/private/internal-path')
            for route in ('import_markdown_file', 'import_original_form'):
                with self.subTest(known=known, route=route), patch.object(
                    form_entries, 'import_document', side_effect=error
                ), patch.object(form_entries, '_save_original_file') as save, patch.object(
                    form_entries.MongoClientManager, 'get_form_templates_collection', return_value=collection
                ):
                    args = {'file': UploadFile(filename='test.hwp', file=io.BytesIO(b'fixture')), 'current_user': user}
                    if route == 'import_original_form':
                        args['template_id'] = str(ObjectId())
                    with self.assertRaises(HTTPException) as failure:
                        await getattr(form_entries, route)(**args)
                    self.assertEqual(failure.exception.status_code, 422)
                    if known:
                        self.assertEqual(failure.exception.detail, message)
                    else:
                        self.assertNotIn('/private/', failure.exception.detail)
                    save.assert_not_called()

    async def test_save_update_and_export_share_markdown(self):
        from bson import ObjectId
        from app.routers import form_entries
        from app.models.form_entry import FormEntryCreate, FormEntryPatch
        collection = AsyncMock()
        entry_id = ObjectId()
        collection.insert_one.return_value = SimpleNamespace(inserted_id=entry_id)
        user = SimpleNamespace(full_name='테스트', email='test@example.com')
        data = {'문서 본문': [{'제목': '테스트', '내용': '# 문서\n\n본문', '내용__format': 'markdown'}]}
        data['기본 정보'] = {'작업명': '원본 양식', '작업 일시': '2026-09-08'}
        with tempfile.TemporaryDirectory() as temporary, patch.object(documents, 'UPLOAD_ROOT', Path(temporary)), patch.object(
            form_entries.MongoClientManager, 'get_form_entries_collection', return_value=collection
        ):
            created = await form_entries.create_entry(FormEntryCreate(template_id=str(ObjectId()), data=data), user)
            stored = collection.insert_one.call_args.args[0]
            self.assertEqual(created.data, data)
            self.assertEqual(Path(stored['markdown_file']).read_text(encoding='utf-8'), data['문서 본문'][0]['내용'])
            data['문서 본문'][0]['내용'] += '\n\n추가 설명'
            collection.find_one_and_update.return_value = {**stored, 'data': data, 'version': 2}
            await form_entries.patch_entry(str(entry_id), FormEntryPatch(data=data, version=1), user)
            update = collection.find_one_and_update.call_args.args[1]['$set']
            self.assertEqual(Path(update['markdown_file']).read_text(encoding='utf-8'), data['문서 본문'][0]['내용'])
            self.assertEqual(update['data']['기본 정보'], data['기본 정보'])
            collection.find_one.return_value = {**stored, 'data': data}
            response = await form_entries.export_entry_hwpx(str(entry_id), user)
            reimported, _ = documents.import_document(response.body, 'test.hwpx')
            self.assertIn('추가 설명', reimported)
            hwp_response = await form_entries.export_entry_hwp(str(entry_id), user)
            self.assertEqual(hwp_response.media_type, 'application/x-hwp')
            self.assertTrue(hwp_response.body.startswith(bytes.fromhex('d0cf11e0a1b11ae1')))
            from app.models.form_entry import FormDocumentExport
            for file_format in ('hwp', 'docx'):
                exported = await form_entries.export_document(FormDocumentExport(markdown=data['문서 본문'][0]['내용'], format=file_format), user)
                self.assertIn(f'.{file_format}', exported.headers['content-disposition'])
                self.assertTrue(exported.body.startswith(bytes.fromhex('d0cf11e0a1b11ae1') if file_format == 'hwp' else b'PK'))


if __name__ == '__main__':
    unittest.main()
