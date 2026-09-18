"""Exercise SR download routes and generated workbooks without accessing MongoDB."""
import io
import unittest
import tempfile
from pathlib import Path
from zipfile import ZipFile
from PIL import Image
from copy import deepcopy
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch
from urllib.parse import unquote

from bson import ObjectId
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.db.mongo import MongoClientManager as M
from app.models.user import UserPublic
from app.routers.auth import get_current_user
from app.routers.sr import admin_requests
from app.services.sr import excel_export


class SRExportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.uploads = tempfile.TemporaryDirectory()
        self.addCleanup(self.uploads.cleanup)
        self.upload_root = Path(self.uploads.name)
        (self.upload_root / 'sr').mkdir()
        (self.upload_root / 'pm').mkdir()
        root_patch = patch.object(excel_export, 'UPLOAD_ROOT', self.upload_root)
        root_patch.start()
        self.addCleanup(root_patch.stop)
        self.sr_id = str(ObjectId())
        self.user = UserPublic(id=str(ObjectId()), email='export@example.com', is_admin=True)
        self.doc = {'_id': ObjectId(self.sr_id), 'sr_no': 'SR-2026-0176', 'title': '서버 점검 요청',
                    'status': 'CLOSED', 'request_type': 'SERVER_INFRA', 'priority': 'MEDIUM',
                    'requester_name': '김요청', 'requester_department': '운영팀',
                    'description': '한글 요청 내용\n두 번째 줄', 'process_result': '처리 완료',
                    'requester_id': str(ObjectId()), 'requester_email': 'requester@example.com',
                    'is_urgent': False, 'compliance_related': False, 'created_by': '김요청', 'updated_by': '김처리',
                    'updated_at': datetime(2026, 9, 17, tzinfo=timezone.utc),
                    'created_at': datetime(2026, 9, 17, tzinfo=timezone.utc)}
        def collection(rows):
            cursor = SimpleNamespace(to_list=AsyncMock(return_value=rows))
            cursor.sort = Mock(return_value=cursor)
            return SimpleNamespace(find=Mock(return_value=cursor))
        self.service = collection([self.doc])
        self.service.find_one = AsyncMock(return_value=self.doc)
        self.db = {
            M.SERVICE_REQUESTS: self.service,
            M.SR_COMMENTS: collection([{'writer_name': '김처리', 'content': '댓글 내용', 'is_internal': True}]),
            M.SR_STATUS_HISTORIES: collection([{'previous_status': 'IN_PROGRESS', 'new_status': 'CLOSED', 'reason': '확인 완료'}]),
            M.SR_HISTORIES: collection([{'action_type': 'FIELD_CHANGE:title', 'before_value': '이전 제목', 'after_value': '서버 점검 요청'}]),
        }
        self.db_patch = patch.object(M, 'get_db', return_value=self.db)
        self.db_patch.start()
        self.addCleanup(self.db_patch.stop)
        app = FastAPI()
        app.include_router(admin_requests.router, prefix='/admin/schedule/service-requests')
        app.dependency_overrides[get_current_user] = lambda: self.user
        self.client = AsyncClient(transport=ASGITransport(app=app, raise_app_exceptions=False), base_url='http://test')
        self.addAsyncCleanup(self.client.aclose)

    async def download(self, path):
        response = await self.client.get('/admin/schedule/service-requests' + path)
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else '')
        self.assertEqual(response.headers['content-type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        disposition = response.headers['content-disposition']
        self.assertTrue(disposition.isascii())
        filename = unquote(disposition.split("filename*=UTF-8''", 1)[1])
        workbook = load_workbook(io.BytesIO(response.content))
        self.addCleanup(workbook.close)
        return filename, workbook

    @staticmethod
    def text(sheet):
        return '\n'.join(str(cell.value) for row in sheet for cell in row if cell.value is not None)

    def file(self, name, content=b'original document', image=False, folder='sr'):
        stored = str(ObjectId()) + Path(name).suffix
        path = self.upload_root / folder / stored
        if image:
            Image.new('RGB', (1200, 700), '#356595').save(path)
        else:
            path.write_bytes(content)
        return {'file_id': stored, 'original_name': name, 'url': f'/api/uploads/{folder}/{stored}',
                'size': path.stat().st_size, 'content_type': 'image/png' if image else 'application/octet-stream'}

    async def bundle(self):
        response = await self.client.get(f'/admin/schedule/service-requests/{self.sr_id}/export')
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else '')
        self.assertEqual(response.headers['content-type'], 'application/zip')
        archive = ZipFile(io.BytesIO(response.content))
        self.addCleanup(archive.close)
        workbook_name = next(name for name in archive.namelist() if name.endswith('.xlsx') and '/' not in name)
        workbook = load_workbook(io.BytesIO(archive.read(workbook_name)))
        self.addCleanup(workbook.close)
        return response, archive, workbook

    async def test_detail_export_readable_sections_and_all_original_content(self):
        original = deepcopy(self.doc)
        filename, workbook = await self.download(f'/{self.sr_id}/export')
        self.assertEqual(filename, 'SR상세_SR-2026-0176.xlsx')
        self.assertEqual(workbook.sheetnames, ['요청정보', '상세내용', '댓글', '상태이력', '필드변경이력', '첨부파일'])
        info = self.text(workbook['요청정보'])
        self.assertIn(self.doc['title'], info)
        self.assertIn(self.doc['description'], info)
        self.assertIn('댓글 내용', self.text(workbook['댓글']))
        self.assertIn('내부 메모', self.text(workbook['댓글']))
        self.assertIn('확인 완료', self.text(workbook['상태이력']))
        self.assertIn('이전 제목', self.text(workbook['필드변경이력']))
        self.assertIn('2026-09-17 09:00', info)
        for sheet in workbook:
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertEqual(sheet.freeze_panes, 'A5')
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertTrue(sheet.print_area)
            self.assertTrue(sheet.merged_cells.ranges)
        self.assertEqual(self.doc, original)

    async def test_images_and_request_comment_originals_are_in_one_portable_bundle(self):
        image = self.file('점검 화면.png', image=True)
        first = self.file('작업 자료.pdf', b'first PDF')
        second = self.file('작업 자료.pdf', b'second PDF')
        self.doc['attachments'] = [image, first]
        comments = self.db[M.SR_COMMENTS].find.return_value.to_list.return_value
        comments[0]['attachments'] = [second, image]
        self.doc['description'] = f"# 작업 요청\n내용 **확인**\n![점검]({image['url']})"
        response, archive, workbook = await self.bundle()
        files = [name for name in archive.namelist() if name.startswith('첨부파일/')]
        self.assertEqual(len(files), 3)
        self.assertEqual(len(set(files)), 3)
        self.assertIn(b'first PDF', [archive.read(name) for name in files])
        self.assertIn(b'second PDF', [archive.read(name) for name in files])
        self.assertEqual(response.headers['x-export-warnings'], '0')
        self.assertEqual(len(workbook['첨부파일']._images), 1)
        self.assertIn('댓글 1', self.text(workbook['첨부파일']))
        self.assertIn('내용 확인', self.text(workbook['요청정보']))
        self.assertNotIn('![점검]', self.text(workbook['요청정보']))
        links = [c.hyperlink.target for row in workbook['첨부파일'] for c in row if c.hyperlink]
        self.assertEqual(set(links), set(files))

    async def test_inline_pm_html_images_and_gif_webp_are_embedded(self):
        inline = self.file('본문.png', image=True, folder='pm')
        self.doc['description'] = f'<p>본문 설명<br>다음 줄</p><img src="{inline["url"]}">'
        attachments = []
        for extension in ('gif', 'webp'):
            item = self.file(f'화면.{extension}', b'')
            Image.new('RGB', (90, 45), 'blue').save(self.upload_root / 'sr' / item['file_id'])
            item['content_type'] = 'image/' + extension
            attachments.append(item)
        self.doc['attachments'] = attachments
        _, archive, workbook = await self.bundle()
        self.assertEqual(len(workbook['첨부파일']._images), 3)
        self.assertEqual(len(archive.namelist()), 4)
        self.assertNotIn('<p>', self.text(workbook['요청정보']))
        self.assertIn('본문 설명\n다음 줄', self.text(workbook['요청정보']))

    async def test_missing_broken_and_external_images_are_reported_without_losing_export(self):
        broken = self.file('깨진 이미지.png', b'not an image')
        broken['content_type'] = 'image/png'
        self.doc['attachments'] = [broken, {'file_id': 'missing.pdf', 'original_name': '없는 문서.pdf'}]
        self.doc['description'] = '![외부](https://example.invalid/private.png)'
        response, archive, workbook = await self.bundle()
        self.assertEqual(response.headers['x-export-warnings'], '3')
        self.assertIn(b'not an image', [archive.read(name) for name in archive.namelist()])
        text = self.text(workbook['첨부파일'])
        self.assertIn('원본 파일을 찾을 수 없습니다', text)
        self.assertIn('이미지 미리보기를 만들 수 없습니다', text)
        self.assertIn('외부 이미지는 포함되지 않았습니다', text)
        self.assertIn('처리 완료', self.text(workbook['요청정보']))

    async def test_type_details_tables_long_text_and_formula_input_are_preserved(self):
        self.doc['title'] = '=HYPERLINK("https://example.invalid", "do not run")'
        self.doc['description'] = '한글 설명과 줄바꿈을 모두 보존합니다. ' * 2000
        self.doc['type_detail'] = {'targetServer': 'server-01', 'workType': 'restart',
                                  'workDetail': '**재기동** 후 서비스 확인',
                                  'firewallRules': [{'sourceIp': '10.0.0.1', 'destinationIp': '10.0.0.2', 'portProtocol': 'TCP/443'}]}
        _, workbook = await self.download(f'/{self.sr_id}/export')
        text = self.text(workbook['상세내용'])
        for expected in ('대상 서버 / 시스템', '서버 재기동', '재기동 후 서비스 확인', '출발지 IP / 대역', 'TCP/443'):
            self.assertIn(expected, text)
        self.assertIn(self.doc['description'].replace('\n', '').strip(), self.text(workbook['요청정보']).replace('\n', ''))
        for sheet in workbook:
            for row in sheet:
                self.assertTrue(all(cell.data_type != 'f' for cell in row))
            self.assertTrue(all((dimension.height or 0) <= 409 for dimension in sheet.row_dimensions.values()))

    async def test_attachment_paths_stay_inside_uploads_and_archive_names_are_safe(self):
        outside = self.upload_root.parent / (self.upload_root.name + '-secret.txt')
        outside.write_text('must not export')
        self.addCleanup(outside.unlink)
        (self.upload_root / 'sr' / 'escape.txt').symlink_to(outside)
        safe = self.file('../../동일\\파일.pdf', b'safe original')
        self.doc['attachments'] = [safe,
            {'file_id': '../' + outside.name, 'original_name': '../../secret.txt'},
            {'file_id': 'escape.txt', 'original_name': 'symlink.txt'}]
        response, archive, _ = await self.bundle()
        self.assertEqual(response.headers['x-export-warnings'], '2')
        for name in archive.namelist():
            self.assertFalse(name.startswith('/'))
            self.assertNotIn('..', Path(name).parts)
            self.assertNotIn('\\', name)
            self.assertNotEqual(archive.read(name), b'must not export')

    async def test_list_export_reaches_static_route_and_downloads_workbook(self):
        filename, workbook = await self.download('/export')
        self.assertTrue(filename.startswith('SR목록_'))
        self.assertEqual(workbook.sheetnames, ['SR목록'])
        self.assertEqual(workbook.active['A2'].value, self.doc['sr_no'])
        self.assertEqual(workbook.active['B2'].value, self.doc['title'])
        self.service.find_one.assert_not_awaited()

    async def test_filename_with_unicode_and_delimiters_remains_one_safe_header(self):
        self.doc['sr_no'] = 'SR-한글 "검토"; ✓\r\n예시'
        filename, _ = await self.download(f'/{self.sr_id}/export')
        self.assertEqual(filename, 'SR상세_SR-한글 _검토_; ✓__예시.xlsx')

    async def test_export_permissions_are_unchanged(self):
        self.user.is_admin = False
        self.user.permissions = ['sr_manager']
        for path in ('/export', f'/{self.sr_id}/export'):
            response = await self.client.get('/admin/schedule/service-requests' + path)
            self.assertEqual(response.status_code, 403, response.text)
        self.service.find_one.assert_not_awaited()
        self.service.find.assert_not_called()

    async def test_invalid_or_missing_sr_returns_json_error(self):
        response = await self.client.get('/admin/schedule/service-requests/not-an-id/export')
        self.assertEqual(response.status_code, 400)
        self.service.find_one.return_value = None
        response = await self.client.get(f'/admin/schedule/service-requests/{self.sr_id}/export')
        self.assertEqual(response.status_code, 404)
        self.assertNotIn('content-disposition', response.headers)

    async def test_detail_json_route_still_resolves_after_static_export_route(self):
        response = await self.client.get(f'/admin/schedule/service-requests/{self.sr_id}')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['sr_no'], self.doc['sr_no'])
